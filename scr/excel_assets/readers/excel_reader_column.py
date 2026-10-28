from pathlib import Path
from typing import Any, Dict, Iterable, Iterator, List, Optional, Union

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter


class ExcelColumnReader:
    """
    Класс для чтения Excel (.xlsx/.xlsm) и получения данных по колонкам.

    Параметры:
        file_path: путь к Excel-файлу.
        sheet: имя листа или индекс листа.
               0 — первый лист, 1 — второй и т.д.
               Если None, берётся активный лист.
        has_header: True, если есть строка с заголовками.
        header_row: номер строки с заголовками. Обычно 1.
        data_only: True — возвращать значения формул, если они сохранены в файле.
                   False — возвращать сами формулы.
    """

    def __init__(
        self,
        file_path: Union[str, Path],
        sheet: Optional[Union[str, int]] = None,
        has_header: bool = True,
        header_row: int = 1,
        data_only: bool = True,
    ) -> None:
        self.file_path = Path(file_path)

        if not self.file_path.exists():
            raise FileNotFoundError(f"Файл не найден: {self.file_path}")

        self._wb = load_workbook(self.file_path, data_only=data_only)

        self._has_header = bool(has_header and header_row is not None and header_row > 0)
        self._header_row = header_row if self._has_header else None
        self._headers_cache: Optional[List[Any]] = None

        self._ws = self._select_sheet(sheet)

    # ------------------------------------------------------------------
    # Служебные методы
    # ------------------------------------------------------------------

    def _select_sheet(self, sheet: Optional[Union[str, int]]):
        if sheet is None:
            return self._wb.active

        if isinstance(sheet, int):
            try:
                return self._wb.worksheets[sheet]
            except IndexError as exc:
                raise ValueError(
                    f"Лист с индексом {sheet} не найден. "
                    f"Доступные листы: {self._wb.sheetnames}"
                ) from exc

        if sheet not in self._wb.sheetnames:
            raise ValueError(
                f"Лист '{sheet}' не найден. "
                f"Доступные листы: {self._wb.sheetnames}"
            )

        return self._wb[sheet]

    @staticmethod
    def _normalize(value: Any) -> str:
        """Приводит значение к строке без лишних пробелов."""
        return "" if value is None else str(value).strip()

    # ------------------------------------------------------------------
    # Свойства
    # ------------------------------------------------------------------

    @property
    def sheet_names(self) -> List[str]:
        """Список всех листов в файле."""
        return self._wb.sheetnames

    @property
    def active_sheet(self) -> str:
        """Имя текущего выбранного листа."""
        return self._ws.title

    @property
    def max_column(self) -> int:
        """Количество столбцов на листе."""
        return self._ws.max_column or 0

    @property
    def headers(self) -> List[Any]:
        """
        Возвращает список заголовков.

        Если has_header=False, вернёт пустой список.
        """
        if not self._has_header or self._header_row is None:
            return []

        if self._headers_cache is None:
            row = next(
                self._ws.iter_rows(
                    min_row=self._header_row,
                    max_row=self._header_row,
                ),
                (),
            )
            self._headers_cache = [cell.value for cell in row]

        return list(self._headers_cache)

    # ------------------------------------------------------------------
    # Управление листами
    # ------------------------------------------------------------------

    def select_sheet(self, sheet: Union[str, int, None]) -> None:
        """
        Переключиться на другой лист.

        Пример:
            reader.select_sheet("Sheet2")
            reader.select_sheet(0)
        """
        self._ws = self._select_sheet(sheet)
        self._headers_cache = None

    # ------------------------------------------------------------------
    # Получение индекса колонки
    # ------------------------------------------------------------------

    def column_index(self, column: Union[int, str]) -> int:
        """
        Возвращает номер столбца в Excel: 1, 2, 3...

        column может быть:
            - номером столбца: 1, 2, 3;
            - заголовком: "Имя", "Возраст";
            - буквой столбца: "A", "B", "C".
        """
        if isinstance(column, int):
            idx = column

        elif isinstance(column, str):
            column = column.strip()

            # Сначала ищем по заголовку.
            if self._has_header:
                target = self._normalize(column)

                for i, header in enumerate(self.headers, start=1):
                    if self._normalize(header) == target:
                        return i

            # Если не нашли заголовок, пробуем как букву столбца.
            if column.isdigit():
                idx = int(column)
            else:
                try:
                    idx = column_index_from_string(column.upper())
                except Exception as exc:
                    raise KeyError(
                        f"Колонка '{column}' не найдена ни по заголовку, "
                        f"ни по букве столбца."
                    ) from exc

        else:
            raise TypeError(
                "column должен быть int, строкой с заголовком "
                "или строкой с буквой столбца."
            )

        if idx < 1 or idx > self.max_column:
            raise IndexError(
                f"Столбец с индексом {idx} вне диапазона. "
                f"Максимальный столбец: {self.max_column}."
            )

        return idx

    # ------------------------------------------------------------------
    # Получение одной колонки
    # ------------------------------------------------------------------

    def get_column(
        self,
        column: Union[int, str],
        include_header: bool = False,
    ) -> List[Any]:
        """
        Вернуть одну колонку как список значений.

        Пример:
            reader.get_column("Имя")
            reader.get_column(1)
            reader.get_column("A")

        Если include_header=True, первая строка с заголовком тоже попадёт в результат.
        """
        idx = self.column_index(column)

        if self._has_header and self._header_row is not None:
            min_row = self._header_row if include_header else self._header_row + 1
        else:
            min_row = 1

        return [
            row[0].value
            for row in self._ws.iter_rows(
                min_row=min_row,
                min_col=idx,
                max_col=idx,
            )
        ]

    # ------------------------------------------------------------------
    # Получение нескольких колонок
    # ------------------------------------------------------------------

    def get_columns(
        self,
        columns: Optional[Iterable[Union[int, str]]] = None,
        include_header: bool = False,
    ) -> Dict[str, List[Any]]:
        """
        Вернуть несколько колонок как словарь.

        Если columns не указан, возвращает все колонки.

        Пример:
            reader.get_columns(["Имя", "Возраст"])
            reader.get_columns([1, 3])
            reader.get_columns(["A", "C"])
        """
        if columns is None:
            if self._has_header:
                headers = self.headers

                if not headers:
                    headers = [None] * self.max_column

                columns = [
                    header
                    if header is not None and self._normalize(header) != ""
                    else get_column_letter(i)
                    for i, header in enumerate(headers, start=1)
                ]
            else:
                columns = list(range(1, self.max_column + 1))

        result: Dict[str, List[Any]] = {}
        used_keys: Dict[str, int] = {}

        for column in columns:
            idx = self.column_index(column)

            if self._has_header and idx <= len(self.headers):
                key = self.headers[idx - 1]

                if key is None or self._normalize(key) == "":
                    key = get_column_letter(idx)
            else:
                key = get_column_letter(idx)

            base_key = str(key)

            # Если есть одинаковые заголовки, делаем ключи уникальными.
            if base_key in used_keys:
                used_keys[base_key] += 1
                unique_key = f"{base_key}_{used_keys[base_key]}"
            else:
                used_keys[base_key] = 1
                unique_key = base_key

            result[unique_key] = self.get_column(idx, include_header=include_header)

        return result

    # ------------------------------------------------------------------
    # Итерация по строкам
    # ------------------------------------------------------------------

    def iter_rows(self, values_only: bool = True):
        """
        Итерация по строкам данных, без строки заголовков.

        Пример:
            for row in reader.iter_rows():
                print(row)
        """
        if self._has_header and self._header_row is not None:
            min_row = self._header_row + 1
        else:
            min_row = 1

        yield from self._ws.iter_rows(min_row=min_row, values_only=values_only)

    # ------------------------------------------------------------------
    # Получение строк как словарей
    # ------------------------------------------------------------------

    def to_records(self, skip_empty_rows: bool = True) -> Iterator[Dict[str, Any]]:
        """
        Вернуть данные как список/итератор словарей.

        Пример:
            for record in reader.to_records():
                print(record["Имя"], record["Возраст"])
        """
        if not self._has_header:
            raise ValueError(
                "to_records работает только если заданы заголовки: has_header=True."
            )

        raw_headers = self.headers

        if not raw_headers:
            raw_headers = [None] * self.max_column

        prepared_headers: List[str] = []
        used_headers: Dict[str, int] = {}

        for i, header in enumerate(raw_headers, start=1):
            name = self._normalize(header)

            if name == "":
                name = f"column_{get_column_letter(i)}"

            if name in used_headers:
                used_headers[name] += 1
                name = f"{name}_{used_headers[name]}"
            else:
                used_headers[name] = 1

            prepared_headers.append(name)

        for row in self.iter_rows(values_only=True):
            if skip_empty_rows and all(value is None for value in row):
                continue

            yield dict(zip(prepared_headers, row))

    # ------------------------------------------------------------------
    # Удобный доступ через []
    # ------------------------------------------------------------------

    def __getitem__(self, column: Union[int, str]) -> List[Any]:
        """
        Позволяет делать так:

            reader["Имя"]
            reader[1]
            reader["A"]
        """
        return self.get_column(column)

    # ------------------------------------------------------------------
    # Закрытие файла / context manager
    # ------------------------------------------------------------------

    def close(self) -> None:
        self._wb.close()

    def __enter__(self) -> "ExcelColumnReader":
        return self

    def __exit__(self, exc_type, exc_val, exc_tb) -> None:
        self.close()

    def __repr__(self) -> str:
        return (
            f"ExcelColumnReader("
            f"file={str(self.file_path)!r}, "
            f"sheet={self.active_sheet!r}, "
            f"max_column={self.max_column}"
            f")"
        )





if __name__=='__main__':
    app = ExcelColumnReader(
        file_path=r"C:\Users\yakovlev_nd\Desktop\Tests\Bam manager\rprd00067mod лтия.464641.003 антенна укв 30-80.xlsx"
    )
    print(app.sheet_names)
    print(app.headers)
    print(app.get_columns([1,2,3,4]))